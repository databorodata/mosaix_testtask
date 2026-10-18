"""Файл ручной проверки карточек: строка базы на строку таблицы, рядом карточки двух моделей,
проверка кодом, привязка к КСР и пустые колонки для оценки.

Запуск: python3 -m scripts.make_review --prep data/debug/prep_50.jsonl --out data/debug/review_50.xlsx
"""

from __future__ import annotations

import argparse
import json

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app import config
from app.nametype import NAMES

HELP = [
    "Как оценивать",
    "ок: карточка верна, числа и коды на месте, лишних фактов нет.",
    "частично: в целом верно, но потеряно важное или синонимы слабые.",
    "ошибка: неверное число, чужая модель или размер, выдуманный факт, неверный тип изделия.",
    "Колонка «проверка кодом»: ok, warn (не всё перенесено), fail (число не из входа или потерян код), error (нет карточки).",
    "КСР: code значит «тот же ресурс», group значит «та же группа», none значит «в КСР нет», service значит услуга. Если КСР привязан неверно, напишите это в комментарии.",
    "Комментарий: что именно не так; это пойдёт в журнал проблем и в правку промпта.",
]


def fmt_card(card: dict | None) -> tuple[str, str, str, str]:
    if not card:
        return "", "", "", ""
    attrs = "; ".join(f"{a['name']}: {a['value']} {a['unit']}".strip() for a in card.get("attributes", []))
    return card.get("canonical_name", ""), card.get("description", ""), "; ".join(card.get("synonyms", [])), attrs


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--prep", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    recs = [json.loads(line) for line in (config.ROOT / args.prep).read_text(encoding="utf-8").splitlines() if line.strip()]
    models = list(recs[0]["cards"].keys())
    a, b = models[0], (models[1] if len(models) > 1 else None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Карточки"
    head = ["№", "Строка базы", "Тип наименования", "Контекст выдачи",
            f"A {a}: название", "A: описание", "A: синонимы", "A: характеристики", "A: проверка кодом"]
    if b:
        head += [f"B {b}: название", "B: описание", "B: проверка кодом"]
    head += ["КСР: статус", "КСР: код", "КСР: наименование", "КСР: единица", "КСР: почему", "Оценка A"]
    if b:
        head += ["Оценка B"]
    head += ["Комментарий"]
    ws.append(head)
    for k, r in enumerate(recs, 1):
        f = r["filter"]
        ctx = f"{f['kind']}, пассажей {len(f['passages'])}; " + ", ".join(sorted({p['domain'] for p in f['passages']}))[:200]
        if f["neighbor_docs"]:
            ctx += f"; отброшены соседи: {len(f['neighbor_docs'])}"
        ca = r["cards"][a]
        va = ca["verify"]
        row = [k, r["name"], NAMES[r["ntype"]], ctx, *fmt_card(ca.get("card")), f"{va['status']}: " + "; ".join(va["issues"])]
        if b:
            cb = r["cards"][b]
            vb = cb["verify"]
            nb, db, _, _ = fmt_card(cb.get("card"))
            row += [nb, db, f"{vb['status']}: " + "; ".join(vb["issues"])]
        kr = r["ksr"]
        row += [kr["status"], kr.get("code", ""), kr.get("name", ""), kr.get("unit", ""), kr.get("reason", ""), ""]
        if b:
            row += [""]
        row += [""]
        ws.append([ILLEGAL_CHARACTERS_RE.sub(" ", v) if isinstance(v, str) else v for v in row])

    widths = {"№": 5, "Строка базы": 40, "Тип наименования": 18, "Контекст выдачи": 26}
    for i, h in enumerate(head, 1):
        col = ws.cell(row=1, column=i).column_letter
        w = widths.get(h, 50 if "описание" in h else 28 if ("синонимы" in h or "характеристики" in h or "наименование" in h) else 16)
        ws.column_dimensions[col].width = w
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="E2ECF9")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"ок,частично,ошибка"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, h in enumerate(head, 1):
        if h.startswith("Оценка"):
            letter = ws.cell(row=1, column=i).column_letter
            dv.add(f"{letter}2:{letter}{len(recs) + 1}")
    hs = wb.create_sheet("Как оценивать")
    for line in HELP:
        hs.append([line])
    hs.column_dimensions["A"].width = 120
    out = config.ROOT / args.out
    wb.save(out)
    print(f"файл проверки: {out.relative_to(config.ROOT)}")


if __name__ == "__main__":
    main()
